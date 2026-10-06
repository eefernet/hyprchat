"""Local Office operations. Deployed with document_formats.py to Codebox.

Edits patch selected XML parts; every other ZIP member is copied byte-for-byte.
Python Office libraries are used to compose new files, never to round-trip an
arbitrary uploaded workbook. LibreOffice only touches conversion/preview copies.
"""
from __future__ import annotations

import copy
import csv
import hashlib
import io
import json
import math
import re
import subprocess
import zipfile
from pathlib import Path

from lxml import etree as ET

from document_formats import (NS, OFFICE_EXTS, LEGACY_EXTS, package, inspect_document,
                              extract_text, slide_paths, sheet_paths, cell_coords, range_bounds)


def parse(raw):
    return ET.fromstring(raw, ET.XMLParser(resolve_entities=False, no_network=True))


def q(prefix, name):
    return "{" + NS[prefix] + "}" + name


def child(node, tag):
    found = node.find(tag)
    return found if found is not None else ET.SubElement(node, tag)


def number(value, lo, hi):
    value = float(value)
    if not math.isfinite(value) or not lo <= value <= hi:
        raise ValueError(f"Number must be between {lo} and {hi}")
    return value


def color(value):
    value = str(value).lstrip("#").upper()
    if not re.fullmatch(r"[A-F0-9]{6}", value):
        raise ValueError("Color must contain six hexadecimal digits")
    return value


def replace_range(nodes, start, end, replacement):
    """Replace a character range without discarding runs, links, or drawings."""
    cursor, inserted = 0, False
    for node in nodes:
        original = node.text or ""
        stop = cursor + len(original)
        if stop > start and cursor < end or (start == end and cursor <= start <= stop and not inserted):
            left, right = max(0, start - cursor), min(len(original), end - cursor)
            node.text = original[:left] + (replacement if not inserted else "") + original[max(left, right):]
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            inserted = True
        cursor = stop
    if not inserted:
        if not nodes:
            raise ValueError("Target has no editable text run")
        nodes[-1].text = (nodes[-1].text or "") + replacement


def edit_text(node, prefix, op):
    # Reject complex field/tracked-change targets, rather than changing their meaning.
    if prefix == "w" and any(node.findall(".//w:" + tag, NS) for tag in ("fldChar", "instrText", "del", "ins")):
        raise ValueError("Fields and tracked-change targets cannot be edited safely")
    nodes = node.findall(".//" + prefix + ":t", NS)
    if op["op"] == "set_text":
        if not nodes:
            run = ET.SubElement(node, q(prefix, "r"))
            nodes = [ET.SubElement(run, q(prefix, "t"))]
        replace_range(nodes, 0, sum(len(n.text or "") for n in nodes), str(op.get("text", "")))
    else:
        source = "".join(n.text or "" for n in nodes)
        find = op.get("find", "")
        if not find or find not in source:
            raise ValueError("Replacement text not found in target")
        matches = list(re.finditer(re.escape(find), source))
        for m in reversed(matches):
            replace_range(nodes, m.start(), m.end(), str(op.get("text", "")))


def format_text(node, prefix, style):
    allowed = {"bold", "italic", "size", "font", "color", "align"}
    if set(style) - allowed:
        raise ValueError("Unsupported text style property")
    for run in node.findall(".//" + prefix + ":r", NS):
        props = run.find(prefix + ":rPr", NS)
        if props is None:
            props = ET.Element(q(prefix, "rPr"))
            run.insert(0, props)
        for key in ("bold", "italic"):
            if key in style:
                attr = "b" if key == "bold" else "i"
                if prefix == "w":
                    child(props, q("w", attr)).set(q("w", "val"), "1" if style[key] else "0")
                else:
                    props.set(attr, "1" if style[key] else "0")
        if "size" in style:
            size = number(style["size"], 1, 400)
            if prefix == "w":
                child(props, q("w", "sz")).set(q("w", "val"), str(round(size * 2)))
            else:
                props.set("sz", str(round(size * 100)))
        if "font" in style:
            if prefix == "w":
                fonts = child(props, q("w", "rFonts"))
                fonts.set(q("w", "ascii"), str(style["font"]))
                fonts.set(q("w", "hAnsi"), str(style["font"]))
            else:
                child(props, q("a", "latin")).set("typeface", str(style["font"]))
        if "color" in style:
            if prefix == "w":
                child(props, q("w", "color")).set(q("w", "val"), color(style["color"]))
            else:
                fill = child(props, q("a", "solidFill"))
                for c in list(fill):
                    fill.remove(c)
                child(fill, q("a", "srgbClr")).set("val", color(style["color"]))
    if "align" in style:
        align = style["align"]
        if align not in {"left", "center", "right", "justify"}:
            raise ValueError("Invalid paragraph alignment")
        props = node.find(prefix + ":pPr", NS)
        if props is None:
            props = ET.Element(q(prefix, "pPr"))
            node.insert(0, props)
        if prefix == "w":
            child(props, q("w", "jc")).set(q("w", "val"), align if align != "justify" else "both")
        else:
            props.set("algn", {"left": "l", "center": "ctr", "right": "r", "justify": "just"}[align])


def find_cell(root, address):
    _, rownum = cell_coords(address)
    data = child(root, q("s", "sheetData"))
    row = data.find(f's:row[@r="{rownum}"]', NS)
    if row is None:
        row = ET.Element(q("s", "row"), r=str(rownum))
        index = next((i for i, r in enumerate(data) if int(r.get("r", 0)) > rownum), len(data))
        data.insert(index, row)
    cell = row.find(f's:c[@r="{address}"]', NS)
    if cell is None:
        cell = ET.Element(q("s", "c"), r=address)
        index = next((i for i, c in enumerate(row) if cell_coords(c.get("r"))[0] > cell_coords(address)[0]), len(row))
        row.insert(index, cell)
    # Keep an explicit used range for streaming readers (which cannot infer it
    # from a missing dimension until a full worksheet scan).
    dimension = root.find("s:dimension", NS)
    if dimension is None:
        dimension = ET.Element(q("s", "dimension"), ref=address)
        root.insert(1 if root.find("s:sheetPr", NS) is not None else 0, dimension)
    c1, r1, c2, r2 = range_bounds(dimension.get("ref", address))
    col, rownum = cell_coords(address)
    def column_name(value):
        letters = ""
        while value:
            value, rem = divmod(value - 1, 26)
            letters = chr(65 + rem) + letters
        return letters
    dimension.set("ref", f"{column_name(min(c1, col))}{min(r1, rownum)}:{column_name(max(c2, col))}{max(r2, rownum)}")
    return cell


def set_cell(root, address, op):
    col, row = cell_coords(address)
    for existing in root.findall(".//s:f", NS):
        if existing.get("t") in {"array", "shared", "dataTable"} and existing.get("ref"):
            c1, r1, c2, r2 = range_bounds(existing.get("ref"))
            if c1 <= col <= c2 and r1 <= row <= r2:
                raise ValueError("Array/shared/data-table formula ranges require an unsupported structural edit")
    cell = find_cell(root, address)
    formula = cell.find("s:f", NS)
    if formula is not None and formula.get("t") in {"array", "shared", "dataTable"}:
        raise ValueError("Array/shared/data-table formulas require an unsupported structural edit")
    for tag in ("v", "is", "f"):
        old = cell.find("s:" + tag, NS)
        if old is not None:
            cell.remove(old)
    cell.attrib.pop("t", None)
    if "formula" in op:
        formula = str(op["formula"]).lstrip("=")
        if not formula or len(formula) > 8192:
            raise ValueError("Invalid formula")
        ET.SubElement(cell, q("s", "f")).text = formula
    elif op.get("value") is not None:
        value = op["value"]
        if isinstance(value, bool):
            cell.set("t", "b")
            ET.SubElement(cell, q("s", "v")).text = "1" if value else "0"
        elif isinstance(value, (int, float)):
            if not math.isfinite(value):
                raise ValueError("Non-finite cell value")
            ET.SubElement(cell, q("s", "v")).text = str(value)
        elif isinstance(value, str):
            if len(value) > 32767:
                raise ValueError("Cell text exceeds Excel limits")
            cell.set("t", "inlineStr")
            text = ET.SubElement(ET.SubElement(cell, q("s", "is")), q("s", "t"))
            text.text = value
            text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        else:
            raise ValueError("Cell values must be strings, numbers, booleans, or null")


def cell_style(parts, cell, style):
    if set(style) - {"bold", "italic", "size", "font", "color", "fill", "number_format", "align", "wrap"}:
        raise ValueError("Unsupported cell style property")
    root = parse(parts["xl/styles.xml"])
    xfs = root.find("s:cellXfs", NS)
    xf = copy.deepcopy(xfs[int(cell.get("s", "0"))])
    if any(k in style for k in ("bold", "italic", "size", "font", "color")):
        fonts = root.find("s:fonts", NS)
        font = copy.deepcopy(fonts[int(xf.get("fontId", "0"))])
        for key, tag in (("bold", "b"), ("italic", "i")):
            if key in style:
                child(font, q("s", tag)).set("val", "1" if style[key] else "0")
        if "size" in style:
            child(font, q("s", "sz")).set("val", str(number(style["size"], 1, 400)))
        if "font" in style:
            child(font, q("s", "name")).set("val", str(style["font"]))
        if "color" in style:
            c = child(font, q("s", "color"))
            c.attrib.clear()
            c.set("rgb", "FF" + color(style["color"]))
        xf.set("fontId", str(len(fonts)))
        xf.set("applyFont", "1")
        fonts.append(font)
        fonts.set("count", str(len(fonts)))
    if "fill" in style:
        fills = root.find("s:fills", NS)
        fill = ET.SubElement(fills, q("s", "fill"))
        pattern = ET.SubElement(fill, q("s", "patternFill"), patternType="solid")
        ET.SubElement(pattern, q("s", "fgColor"), rgb="FF" + color(style["fill"]))
        xf.set("fillId", str(len(fills) - 1))
        xf.set("applyFill", "1")
        fills.set("count", str(len(fills)))
    if "number_format" in style:
        formats = root.find("s:numFmts", NS)
        if formats is None:
            formats = ET.Element(q("s", "numFmts"))
            root.insert(0, formats)
        num = max([163] + [int(n.get("numFmtId")) for n in formats]) + 1
        ET.SubElement(formats, q("s", "numFmt"), numFmtId=str(num), formatCode=str(style["number_format"]))
        formats.set("count", str(len(formats)))
        xf.set("numFmtId", str(num))
        xf.set("applyNumberFormat", "1")
    if "align" in style or "wrap" in style:
        align = child(xf, q("s", "alignment"))
        if "align" in style:
            if style["align"] not in {"left", "center", "right", "justify"}:
                raise ValueError("Invalid cell alignment")
            align.set("horizontal", style["align"])
        if "wrap" in style:
            align.set("wrapText", "1" if style["wrap"] else "0")
        xf.set("applyAlignment", "1")
    cell.set("s", str(len(xfs)))
    xfs.append(xf)
    xfs.set("count", str(len(xfs)))
    parts["xl/styles.xml"] = ET.tostring(root, xml_declaration=True, encoding="UTF-8")


def edit_document(source, destination, operations, assets=None):
    parts = package(source)
    original = dict(parts)
    if any(n.startswith("_xmlsignatures/") for n in parts):
        raise ValueError("Editing would invalidate the document's digital signature")
    if not isinstance(operations, list) or not 1 <= len(operations) <= 1000:
        raise ValueError("Provide between 1 and 1000 operations")
    roots = {}
    for op in operations:
        action = op.get("op")
        if action == "set_chart_data":
            target = op.get("target", "")
            if not target.startswith("ppt/charts/") or target not in parts:
                raise ValueError("Chart-data edits support PowerPoint category charts; edit worksheet cells for Excel charts")
            from pptx import Presentation
            from pptx.chart.data import CategoryChartData
            from document_formats import relationships
            chart_root = parse(parts[target])
            embeddings = {n for n in relationships(parts, target).values() if n.startswith("ppt/embeddings/")}
            if len(embeddings) != 1 or not embeddings.issubset(parts):
                raise ValueError("Chart edits require one embedded workbook; externally linked charts are unsupported")
            if not any(chart_root.findall(".//c:" + kind, NS) for kind in ("barChart", "lineChart", "pieChart")):
                raise ValueError("Only bar, line, and pie chart data can be replaced")
            data = CategoryChartData()
            data.categories = op["categories"]
            for series in op["series"]:
                if len(series["values"]) != len(op["categories"]):
                    raise ValueError("Chart series lengths must match categories")
                data.add_series(series["name"], series["values"])
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                for name, raw in parts.items():
                    z.writestr(name, raw)
            buf.seek(0)
            deck = Presentation(buf)
            chart = next((s.chart for slide in deck.slides for s in slide.shapes
                          if s.has_chart and str(s.chart.part.partname).lstrip("/") == target), None)
            if chart is None:
                raise ValueError("Chart target not found")
            chart.replace_data(data)
            updated = io.BytesIO()
            deck.save(updated)
            changed = {target} | embeddings
            with zipfile.ZipFile(updated) as z:
                for name in changed:
                    parts[name] = z.read(name)
            continue
        if action == "reorder_slides":
            paths = slide_paths(parts)
            order = op.get("order", [])
            if sorted(order) != list(range(1, len(paths) + 1)):
                raise ValueError("Slide order must contain every slide number exactly once")
            root = roots.setdefault("ppt/presentation.xml", parse(parts["ppt/presentation.xml"]))
            listing = root.find("p:sldIdLst", NS)
            slides = list(listing)
            listing[:] = [slides[i - 1] for i in order]
            continue
        if action == "replace_image":
            target = op.get("target", "")
            if target not in parts or "/media/" not in target:
                raise ValueError("Image target not found")
            asset = (assets or {}).get(op.get("image_artifact_id"))
            if not asset:
                raise ValueError("Image artifact not found")
            from PIL import Image
            with Image.open(asset) as im:
                fmt = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG"}.get(Path(target).suffix.lower())
                if not fmt:
                    raise ValueError("Only PNG/JPEG image targets can be replaced")
                buffer = io.BytesIO()
                im.convert("RGB" if fmt == "JPEG" else "RGBA").save(buffer, format=fmt)
                parts[target] = buffer.getvalue()
            continue
        if action == "page_setup":
            root = roots.setdefault("word/document.xml", parse(parts["word/document.xml"]))
            sections = root.findall(".//w:sectPr", NS)
            for section in sections:
                for key, tag, attr in (("width", "pgSz", "w"), ("height", "pgSz", "h"),
                                       ("margin_top", "pgMar", "top"), ("margin_bottom", "pgMar", "bottom"),
                                       ("margin_left", "pgMar", "left"), ("margin_right", "pgMar", "right")):
                    if key in op:
                        child(section, q("w", tag)).set(q("w", attr), str(round(number(op[key], 0.1, 30) * 1440)))
            continue
        target = op.get("target", "")
        name, sep, selector = target.partition("#")
        if not sep or name not in parts or not name.endswith(".xml"):
            raise ValueError("Target not found; read the source document first")
        root = roots.setdefault(name, parse(parts[name]))
        if selector.startswith("shape:") and name.startswith("ppt/slides/") and action == "format_shape":
            shape = next((s for s in root.findall("p:cSld/p:spTree/*", NS)
                          if s.find(".//p:cNvPr", NS) is not None and s.find(".//p:cNvPr", NS).get("id") == selector[6:]), None)
            if shape is None:
                raise ValueError("Slide shape not found")
            transform = shape.find(".//a:xfrm", NS)
            if transform is None:
                transform = shape.find("p:xfrm", NS)
            if transform is None:
                raise ValueError("Shape inherits its layout; explicit geometry is unavailable")
            for key, tag, attr in (("x", "off", "x"), ("y", "off", "y"), ("width", "ext", "cx"), ("height", "ext", "cy")):
                if key in op:
                    child(transform, q("a", tag)).set(attr, str(round(number(op[key], 0 if key in {"x", "y"} else .01, 30) * 914400)))
        elif selector.startswith("p:") and action in {"set_text", "replace_text", "format_text"}:
            prefix = "w" if name.startswith("word/") else "a" if name.startswith("ppt/") else ""
            if not prefix:
                raise ValueError("Not a paragraph target")
            paragraphs = root.findall(".//" + prefix + ":p", NS)
            index = int(selector[2:])
            if index < 0 or index >= len(paragraphs):
                raise ValueError("Paragraph target not found")
            node = paragraphs[index]
            if action == "format_text":
                format_text(node, prefix, op.get("style", {}))
            else:
                edit_text(node, prefix, op)
        elif selector.startswith("cell:") and name.startswith("xl/worksheets/") and action in {"set_cell", "format_cell"}:
            address = selector[5:].upper()
            cell_coords(address)
            if root.find("s:sheetProtection", NS) is not None:
                raise ValueError("Protected worksheets cannot be edited")
            if action == "set_cell":
                set_cell(root, address, op)
            else:
                cell_style(parts, find_cell(root, address), op.get("style", {}))
        else:
            raise ValueError("Unsupported operation/target combination")
    if str(source).endswith(".xlsx") and any(o.get("op") == "set_cell" for o in operations):
        root = roots.setdefault("xl/workbook.xml", parse(parts["xl/workbook.xml"]))
        calc = child(root, q("s", "calcPr"))
        calc.set("fullCalcOnLoad", "1")
        calc.set("forceFullCalc", "1")
        # A dependency chain can become invalid after changing formulas. Conservatively
        # block such edits until the caller normalizes an explicit conversion copy.
        if "xl/calcChain.xml" in parts:
            raise ValueError("Workbook contains a calculation chain; normalize a copy before editing formulas/cells")
    for name, root in roots.items():
        parts[name] = ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as out:
        for name, raw in parts.items():
            out.writestr(name, raw)
    checked = package(destination)
    if set(checked) != set(original):
        raise ValueError("Package preservation check failed")
    changed = [name for name in checked if checked[name] != original[name]]
    return {"changed_parts": changed, "preserved_parts": len(parts) - len(changed),
            "operations_applied": len(operations)}


def _docx_paragraph(parent, spec):
    from docx.shared import Pt, RGBColor
    p = parent.add_paragraph(style=spec.get("style"))
    if "heading" in spec:
        p.style = "Title" if int(spec["heading"]) == 0 else f'Heading {int(number(spec["heading"], 1, 9))}'
    for run in spec.get("runs", [{"text": spec.get("text", "")} ]):
        r = p.add_run(str(run.get("text", "")))
        for key in ("bold", "italic"):
            if key in run:
                setattr(r, key, bool(run[key]))
        if "size" in run:
            r.font.size = Pt(number(run["size"], 1, 400))
        if "color" in run:
            r.font.color.rgb = RGBColor.from_string(color(run["color"]))
    return p


def create_document(destination, content, assets=None):
    if not isinstance(content, dict):
        raise ValueError("content must be an object")
    assets = assets or {}
    ext = Path(destination).suffix
    if ext == ".docx":
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        doc.styles["Normal"].font.name = "Liberation Sans"
        for block in content.get("blocks", []):
            kind = block.get("type", "paragraph")
            if kind == "paragraph":
                _docx_paragraph(doc, block)
            elif kind == "table":
                rows = block["rows"]
                if not rows or not rows[0] or len(rows) * len(rows[0]) > 20000:
                    raise ValueError("Invalid table dimensions")
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    if len(row) != len(rows[0]):
                        raise ValueError("Table rows must have equal lengths")
                    for j, value in enumerate(row):
                        table.cell(i, j).text = str(value)
            elif kind == "image":
                doc.add_picture(assets[block["image_artifact_id"]], width=Inches(number(block.get("width", 5), 0.1, 20)))
                if block.get("caption"):
                    doc.add_paragraph(block["caption"], "Caption")
            elif kind == "page_break":
                doc.add_page_break()
            else:
                raise ValueError("Unsupported Word block")
        for section in doc.sections:
            if "header" in content:
                section.header.paragraphs[0].text = str(content["header"])
            if "footer" in content:
                section.footer.paragraphs[0].text = str(content["footer"])
        doc.save(destination)
    elif ext == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        slides = content.get("slides", [])
        if not 1 <= len(slides) <= 200:
            raise ValueError("Provide between 1 and 200 slides")
        for spec in slides:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            title = slide.shapes.add_textbox(Inches(.6), Inches(.3), Inches(12), Inches(.8)).text_frame
            title.text = spec.get("title", "")
            title.paragraphs[0].font.size = Pt(30)
            title.paragraphs[0].font.name = "Liberation Sans"
            elements = spec.get("elements", [])
            if spec.get("bullets"):
                elements = [{"type": "text", "text": "\n".join(spec["bullets"])}] + elements
            for element in elements:
                x, y, w, h = [Inches(number(element.get(k, default), 0 if k in {"x", "y"} else .1, 30))
                              for k, default in (("x", .7), ("y", 1.5), ("width", 11.8), ("height", 4.8))]
                kind = element.get("type", "text")
                if kind == "text":
                    frame = slide.shapes.add_textbox(x, y, w, h).text_frame
                    frame.word_wrap = True
                    frame.text = str(element.get("text", ""))
                    for p in frame.paragraphs:
                        p.font.size = Pt(number(element.get("size", 22), 1, 400))
                        p.font.name = "Liberation Sans"
                elif kind == "image":
                    slide.shapes.add_picture(assets[element["image_artifact_id"]], x, y, width=w, height=h)
                elif kind == "table":
                    rows = element["rows"]
                    if not rows or not rows[0] or len(rows) * len(rows[0]) > 1000:
                        raise ValueError("Invalid slide table dimensions")
                    table = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, h).table
                    for i, row in enumerate(rows):
                        if len(row) != len(rows[0]):
                            raise ValueError("Table rows must have equal lengths")
                        for j, value in enumerate(row):
                            table.cell(i, j).text = str(value)
                elif kind == "chart":
                    data = CategoryChartData()
                    data.categories = element["categories"]
                    for series in element["series"]:
                        data.add_series(series["name"], series["values"])
                    chart_type = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}[element.get("chart_type", "bar")]
                    slide.shapes.add_chart(chart_type, x, y, w, h, data)
                else:
                    raise ValueError("Unsupported slide element")
            if "notes" in spec:
                slide.notes_slide.notes_text_frame.text = str(spec["notes"])
        deck.save(destination)
    elif ext == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.worksheet.table import Table, TableStyleInfo
        book = Workbook()
        book.remove(book.active)
        specs = content.get("sheets", [])
        if not 1 <= len(specs) <= 100:
            raise ValueError("Provide between 1 and 100 worksheets")
        for spec in specs:
            sh = book.create_sheet(spec["name"])
            for row in spec.get("rows", []):
                sh.append(row)
            for address, value in spec.get("cells", {}).items():
                cell_coords(address)
                sh[address] = value
            for fmt in spec.get("formats", []):
                c1, r1, c2, r2 = range_bounds(fmt["range"])
                if (c2 - c1 + 1) * (r2 - r1 + 1) > 50000:
                    raise ValueError("Formatting range too large")
                style = fmt["style"]
                for row in sh.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                    for cell in row:
                        cell.font = Font(name=style.get("font", "Liberation Sans"), size=number(style.get("size", 11), 1, 400),
                                         bold=bool(style.get("bold")), italic=bool(style.get("italic")), color=color(style.get("color", "000000")))
                        if "fill" in style:
                            cell.fill = PatternFill("solid", fgColor=color(style["fill"]))
                        if "number_format" in style:
                            cell.number_format = style["number_format"]
                        cell.alignment = Alignment(horizontal=style.get("align", "general"), wrap_text=bool(style.get("wrap")))
            if spec.get("freeze_panes"):
                sh.freeze_panes = spec["freeze_panes"]
            for column, width in spec.get("column_widths", {}).items():
                cell_coords(column + "1")
                sh.column_dimensions[column].width = number(width, 1, 255)
            for table in spec.get("tables", []):
                range_bounds(table["range"])
                obj = Table(displayName=table["name"], ref=table["range"])
                obj.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
                sh.add_table(obj)
            for spec_chart in spec.get("charts", []):
                chart = {"bar": BarChart, "line": LineChart, "pie": PieChart}[spec_chart.get("type", "bar")]()
                c1, r1, c2, r2 = range_bounds(spec_chart["range"])
                chart.add_data(Reference(sh, min_col=c1 + 1, min_row=r1, max_col=c2, max_row=r2), titles_from_data=True)
                chart.set_categories(Reference(sh, min_col=c1, min_row=r1 + 1, max_row=r2))
                chart.title = spec_chart.get("title", "")
                sh.add_chart(chart, spec_chart.get("anchor", "H2"))
        book.save(destination)
    else:
        raise ValueError("Create supports docx, pptx, and xlsx")
    package(destination)


def libreoffice(source, outdir, extension):
    """Run inside the worker's networkless sandbox; profile is unique per job."""
    from pathlib import Path
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    profile = outdir.parent / (outdir.name + "-profile")
    profile.mkdir(exist_ok=True)
    # Macro security remains enabled even for legacy files with embedded VBA.
    (profile / "user").mkdir(exist_ok=True)
    (profile / "user" / "registrymodifications.xcu").write_text(
        '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
        '<item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item></oor:items>')
    result = subprocess.run(["libreoffice", "-env:UserInstallation=" + profile.resolve().as_uri(),
                             "--headless", "--norestore", "--convert-to", extension,
                             "--outdir", str(outdir), str(source)], capture_output=True, timeout=120)
    output = outdir / (Path(source).stem + "." + extension.split(":")[0])
    if result.returncode or not output.is_file() or output.stat().st_size == 0:
        raise ValueError("LibreOffice conversion failed")
    return output


def execute(request, workdir):
    workdir = Path(workdir)
    action = request["action"]
    source = workdir / request["source"] if request.get("source") else None
    assets = {k: str(workdir / v) for k, v in request.get("assets", {}).items()}
    warnings, outputs, report = [], [], {}
    if source and source.suffix.lower() in LEGACY_EXTS:
        import olefile
        with olefile.OleFileIO(source) as ole:
            entries = {"/".join(p).lower() for p in ole.listdir()}
            if any("vba" in n or "macros" in n or "encryptedpackage" in n or "encryptioninfo" in n for n in entries):
                raise ValueError("Encrypted or macro-enabled legacy files are unsupported")
        modern = {".doc": "docx", ".ppt": "pptx", ".xls": "xlsx"}[source.suffix.lower()]
        source = libreoffice(source, workdir / "normalized", modern)
        package(source)
        warnings.append("Legacy Office input was normalized using LibreOffice; layout may differ from the original.")
        outputs.append({"path": str(source.relative_to(workdir)), "role": "normalized"})
    if source:
        package(source)
    if action == "read":
        report = inspect_document(source, offset=request.get("offset", 0), limit=request.get("limit", 100),
                                  sheet=request.get("sheet", ""), cell_range=request.get("range", ""))
    elif action in {"create", "edit"}:
        ext = source.suffix if action == "edit" else "." + request["format"]
        output = workdir / ("result" + ext)
        if action == "create":
            create_document(output, request.get("content", {}), assets)
        else:
            report.update(edit_document(source, output, request.get("operations", []), assets))
        source = output
        report["inspection"] = inspect_document(source, limit=100)
        outputs.append({"path": output.name, "role": "document"})
    elif action == "convert":
        fmt = request.get("format", "pdf")
        if fmt in {"txt", "md"} and source.suffix in {".docx", ".pptx"}:
            output = workdir / ("result." + fmt)
            output.write_text(extract_text(source), encoding="utf-8")
            if fmt == "md":
                warnings.append("Markdown export contains extracted text, not a layout-preserving representation.")
        elif fmt == "csv" and source.suffix == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(source, read_only=True, data_only=True)
            if not request.get("sheet") or request["sheet"] not in book.sheetnames:
                raise ValueError("CSV export requires an explicit worksheet name")
            sh = book[request["sheet"]]
            if sh.max_row is None or sh.max_column is None:
                sh.calculate_dimension(force=True)
            if sh.max_row * sh.max_column > 1000000:
                raise ValueError("CSV export exceeds one million cells")
            output = workdir / "result.csv"
            with output.open("w", encoding="utf-8", newline="") as handle:
                csv.writer(handle).writerows(sh.values)
            book.close()
            warnings.append("CSV contains source cached values; formula results may be stale or empty.")
        elif fmt == "pdf":
            output = libreoffice(source, workdir / "export", "pdf")
        elif fmt == source.suffix[1:]:
            output = libreoffice(source, workdir / "export", fmt)
            package(output)
            warnings.append("Re-saved through LibreOffice; advanced features and layout may differ. The original is retained.")
        else:
            raise ValueError("Unsupported conversion for this document type")
        outputs.append({"path": str(output.relative_to(workdir)), "role": "export"})
    elif action != "preview":
        raise ValueError("Unknown document action")
    if action in {"preview", "create", "edit"}:
        if source.suffix == ".xlsx":
            report["inspection"] = inspect_document(source, sheet=request.get("sheet", ""), cell_range=request.get("range", ""), limit=100)
        try:
            preview = libreoffice(source, workdir / "preview", "pdf")
            outputs.append({"path": str(preview.relative_to(workdir)), "role": "preview"})
            warnings.append("Preview rendered by LibreOffice; Microsoft Office pagination may differ.")
        except (ValueError, OSError, subprocess.TimeoutExpired):
            warnings.append("Preview unavailable; the Office file remains downloadable. Layout has not been verified.")
    if source:
        report["text"] = extract_text(source, max_chars=200000)
    return {"report": report, "warnings": warnings, "outputs": outputs}
