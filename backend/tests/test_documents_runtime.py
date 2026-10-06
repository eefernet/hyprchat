"""Behavioral document tests: preservation, independent library reopening, exports."""
import hashlib
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest.importorskip("docx")
pytest.importorskip("pptx")
pytest.importorskip("openpyxl")
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from PIL import Image

from document_formats import inspect_document, extract_text, package, NS
from document_runtime import create_document, edit_document, execute


def test_word_edit_across_runs_preserves_images_and_formatting(tmp_path):
    image = tmp_path / "picture.png"
    Image.new("RGB", (30, 20), "red").save(image)
    source, output = tmp_path / "report.docx", tmp_path / "edited.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("The old ").bold = True
    p.add_run("report is ready.").italic = True
    doc.add_picture(str(image))
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Keep table"
    doc.sections[0].header.paragraphs[0].text = "Preserve header"
    doc.save(source)
    before = source.read_bytes()
    result = inspect_document(source)
    edit_document(source, output, [{"op": "replace_text", "target": result["items"][0]["target"], "find": "old report", "text": "updated summary"}])
    changed = Document(output)
    assert changed.paragraphs[0].text == "The updated summary is ready."
    assert changed.paragraphs[0].runs[0].bold is True
    assert changed.paragraphs[0].runs[1].italic is True
    assert changed.tables[0].cell(0, 0).text == "Keep table"
    assert len(changed.inline_shapes) == 1
    assert source.read_bytes() == before
    original, revised = package(source), package(output)
    assert {k for k in original if original[k] != revised[k]} == {"word/document.xml"}


def test_word_headers_tables_and_page_setup(tmp_path):
    path, output = tmp_path / "new.docx", tmp_path / "revision.docx"
    create_document(path, {"blocks": [{"heading": 1, "text": "Heading"}, {"type": "table", "rows": [["Name", "Value"], ["A", "10"]]}], "header": "Draft"})
    inspection = inspect_document(path)
    target = next(i["target"] for i in inspection["items"] if i["text"] == "Draft")
    edit_document(path, output, [{"op": "set_text", "target": target, "text": "Final"}, {"op": "page_setup", "margin_left": 1.25}])
    doc = Document(output)
    assert doc.sections[0].header.paragraphs[0].text == "Final"
    assert round(doc.sections[0].left_margin.inches, 2) == 1.25
    assert "Name" in extract_text(output)


def test_powerpoint_notes_slide_order_and_charts_preserved(tmp_path):
    path, output = tmp_path / "deck.pptx", tmp_path / "edited.pptx"
    create_document(path, {"slides": [
        {"title": "First", "notes": "Speaker note", "elements": [{"type": "chart", "categories": ["A", "B"], "series": [{"name": "Revenue", "values": [3, 6]}]}]},
        {"title": "Second", "bullets": ["Keep me"]},
    ]})
    before = package(path)
    read = inspect_document(path)
    target = next(i["target"] for i in read["items"] if i.get("text") == "First")
    assert any(i.get("notes") and "Speaker note" in i["text"] for i in read["items"])
    edit_document(path, output, [{"op": "set_text", "target": target, "text": "Revised"}, {"op": "reorder_slides", "order": [2, 1]}])
    deck = Presentation(output)
    assert deck.slides[0].shapes[0].text == "Second"
    assert deck.slides[1].shapes[0].text == "Revised"
    after = package(output)
    assert all(after[k] == v for k, v in before.items() if "/charts/" in k or "/embeddings/" in k)


def test_excel_targeted_edit_preserves_unknown_parts_and_formulas(tmp_path):
    path, output = tmp_path / "book.xlsx", tmp_path / "edited.xlsx"
    create_document(path, {"sheets": [{"name": "Budget", "rows": [["Item", "Cost"], ["A", 20]], "cells": {"B3": "=SUM(B2:B2)"},
                                      "charts": [{"range": "A1:B2", "title": "Costs"}]}, {"name": "Untouched", "rows": [["Keep"]]}]})
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("customXml/kept.xml", '<custom xmlns="urn:hyprchat-test">preserve</custom>')
    inspection = inspect_document(path, sheet="Budget", cell_range="A1:B3")
    target = next(i["target"] for i in inspection["items"] if i["cell"] == "B2")
    assert next(i for i in inspection["items"] if i["cell"] == "B3")["formula"] == "=SUM(B2:B2)"
    edit_document(path, output, [{"op": "set_cell", "target": target, "value": 42}, {"op": "format_cell", "target": target, "style": {"bold": True, "fill": "FFFF00", "number_format": "0.00"}}])
    book = load_workbook(output)
    assert book["Budget"]["B2"].value == 42
    assert book["Budget"]["B3"].value == "=SUM(B2:B2)"
    assert book["Budget"]["B2"].font.bold
    assert book["Untouched"]["A1"].value == "Keep"
    before, after = package(path), package(output)
    assert before["customXml/kept.xml"] == after["customXml/kept.xml"]
    assert before["xl/worksheets/sheet2.xml"] == after["xl/worksheets/sheet2.xml"]
    assert all(after[k] == v for k, v in before.items() if "/charts/" in k)


def test_excel_formula_edit_new_cell_and_literal_equals(tmp_path):
    path, output = tmp_path / "book.xlsx", tmp_path / "edited.xlsx"
    create_document(path, {"sheets": [{"name": "Sheet", "rows": [[1]]}]})
    edit_document(path, output, [{"op": "set_cell", "target": "xl/worksheets/sheet1.xml#cell:C8", "formula": "=SUM(A1:A1)"},
                                 {"op": "set_cell", "target": "xl/worksheets/sheet1.xml#cell:B2", "value": "=literal"}])
    book = load_workbook(output)
    assert book.active["C8"].value == "=SUM(A1:A1)"
    assert book.active["B2"].data_type == "s"
    readonly = load_workbook(output, read_only=True)
    assert (readonly.active.max_row, readonly.active.max_column) == (8, 3)
    readonly.close()
    result = execute({"action": "convert", "source": output.name, "format": "csv", "sheet": "Sheet"}, tmp_path)
    assert result["outputs"][0]["role"] == "export"


def test_read_pagination(tmp_path):
    path = tmp_path / "many.docx"
    create_document(path, {"blocks": [{"text": str(i)} for i in range(5)]})
    first = inspect_document(path, limit=2)
    last = inspect_document(path, offset=4, limit=2)
    assert first["next_offset"] == 2
    assert last["next_offset"] is None
    assert last["items"][0]["text"] == "4"


def test_powerpoint_chart_update_and_geometry_preserve_other_parts(tmp_path):
    path, output = tmp_path / "deck.pptx", tmp_path / "out.pptx"
    create_document(path, {"slides": [{"title": "Chart", "elements": [{"type": "chart", "categories": ["A", "B"], "series": [{"name": "Old", "values": [1, 2]}]}]}]})
    read = inspect_document(path)
    edit_document(path, output, [{"op": "set_chart_data", "target": read["charts"][0]["target"], "categories": ["C", "D"], "series": [{"name": "New", "values": [5, 8]}]},
                                 {"op": "format_shape", "target": read["shapes"][0]["target"], "x": 1, "width": 10}])
    deck = Presentation(output)
    chart = next(s.chart for s in deck.slides[0].shapes if s.has_chart)
    assert list(chart.series[0].values) == [5, 8]
    assert chart.series[0].name == "New"
    assert deck.slides[0].shapes[0].left.inches == 1
    before, after = package(path), package(output)
    assert before["ppt/theme/theme1.xml"] == after["ppt/theme/theme1.xml"]


@pytest.mark.parametrize("member,payload", [("../escape", b"x"), ("word/evil.xml", b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>'), ("word/vbaProject.bin", b"macro")])
def test_reject_unsafe_packages(tmp_path, member, payload):
    path = tmp_path / "bad.docx"
    create_document(path, {"blocks": [{"text": "safe"}]})
    with zipfile.ZipFile(path, "a") as z:
        z.writestr(member, payload)
    with pytest.raises(ValueError):
        inspect_document(path)


def test_failed_edit_leaves_original_unchanged(tmp_path):
    path, output = tmp_path / "book.docx", tmp_path / "out.docx"
    create_document(path, {"blocks": [{"text": "Keep"}]})
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    with pytest.raises(ValueError):
        edit_document(path, output, [{"op": "delete_everything", "target": "word/document.xml#p:0"}])
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    assert not output.exists()


def test_preview_failure_does_not_lose_valid_output(tmp_path, monkeypatch):
    import document_runtime
    def unavailable(*args):
        raise OSError("LibreOffice missing")
    monkeypatch.setattr(document_runtime, "libreoffice", unavailable)
    result = execute({"action": "create", "format": "docx", "content": {"blocks": [{"text": "Ready"}]}}, tmp_path)
    assert result["outputs"] == [{"path": "result.docx", "role": "document"}]
    assert "Preview unavailable" in result["warnings"][0]
    assert Document(tmp_path / "result.docx").paragraphs[0].text == "Ready"


def test_csv_requires_named_sheet_and_retains_cached_warning(tmp_path):
    path = tmp_path / "source.xlsx"
    create_document(path, {"sheets": [{"name": "Data", "rows": [["A", 2]]}]})
    with pytest.raises(ValueError, match="explicit worksheet"):
        execute({"action": "convert", "source": path.name, "format": "csv"}, tmp_path)
    result = execute({"action": "convert", "source": path.name, "format": "csv", "sheet": "Data"}, tmp_path)
    assert (tmp_path / "result.csv").read_text().strip() == "A,2"
    assert any("cached" in w for w in result["warnings"])


def test_excel_array_range_rejects_non_anchor_edit(tmp_path):
    from lxml import etree as ET
    path, output = tmp_path / 'array.xlsx', tmp_path / 'revision.xlsx'
    create_document(path, {'sheets': [{'name': 'Data', 'rows': [[1, 2], [3, 4]]}]})
    parts = package(path)
    root = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    anchor = root.find('.//s:c[@r="A1"]', NS)
    ET.SubElement(anchor, '{'+NS['s']+'}f', t='array', ref='A1:B2').text = 'ROW(A1:B2)'
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(root)
    with zipfile.ZipFile(path, 'w') as z:
        for name, data in parts.items():
            z.writestr(name, data)
    original = path.read_bytes()
    with pytest.raises(ValueError, match='formula ranges'):
        edit_document(path, output, [{'op': 'set_cell', 'target': 'xl/worksheets/sheet1.xml#cell:B2', 'value': 5}])
    assert path.read_bytes() == original
    assert not output.exists()
