"""Workspace hub qualification against isolated SQLite and captured model requests."""
import asyncio
import hashlib
import io
import json
import time
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
import database as db
import workspace_service as ws


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    import config
    monkeypatch.setattr(db,'DATABASE_PATH',str(tmp_path/'hub.db'))
    monkeypatch.setattr(config,'KB_DIR',str(tmp_path/'knowledge'))
    monkeypatch.setattr(config,'ARTIFACT_DIR',str(tmp_path/'artifacts'))
    asyncio.run(db.init_db())
    return tmp_path


def test_defaults_atomic_creation_and_invalid_selection(isolated,monkeypatch):
    import main
    monkeypatch.setattr(main,'list_models',AsyncMock(return_value={'models':['app','persona','workspace','explicit']}))
    async def run():
        await db.create_workspace('w','Project')
        await db.create_model_config('p','Persona','persona','Persona instructions',['research'])
        await db.update_workspace('w',chat_defaults={'model':'workspace','model_config_id':'p','tool_ids':[]})
        created=await main.create_conversation(main.ConversationCreate(workspace_context={'mode':'selected','workspace_id':'w'}))
        assert created['model']=='workspace' and created['model_config_id']=='p' and created['tool_ids']==[]
        assert created['system_prompt']=='Persona instructions' and created['persona_name']=='Persona'
        assert created['effective_workspace']['id']=='w'
        explicit=await main.create_conversation(main.ConversationCreate(model='explicit',tool_ids=['codeagent'],model_config_id='',workspace_context={'mode':'selected','workspace_id':'w'}))
        assert explicit['model']=='explicit' and explicit['model_config_id'] is None and explicit['tool_ids']==['codeagent']
        await db.update_workspace('w',chat_defaults={'model_config_id':'p'})
        assert (await ws.resolve_defaults({},'w'))['tool_ids']==['research']
        assert (await ws.resolve_defaults({},'w'))['model']=='persona'
        await db.update_workspace('w',chat_defaults={'model':'missing'})
        before=len(await db.get_conversations())
        from fastapi import HTTPException
        with pytest.raises(HTTPException,match='') as exc:
            await main.create_conversation(main.ConversationCreate(workspace_context={'mode':'selected','workspace_id':'w'}))
        assert exc.value.status_code==400 and 'Choose an available model' in exc.value.detail
        assert len(await db.get_conversations())==before
        await db.update_workspace('w',chat_defaults={'model_config_id':'missing'})
        with pytest.raises(ValueError,match='persona is unavailable'):await ws.resolve_defaults({},'w')
        saved=await db.get_conversation(created['id'])
        assert saved['model']=='workspace' and saved['tool_ids']==[]
        await db.update_conversation(saved['id'],workspace_context={'mode':'off'})
        assert (await db.get_conversation(saved['id']))['model']=='workspace'
    asyncio.run(run())


def test_migration_preserves_memory_and_instructions_are_independent(isolated):
    async def run():
        await db.create_workspace('w','Project')
        await db.update_workspace('w',instructions='Use metric units.',memory_enabled=0)
        await db.create_workspace_memory('w',content='Accepted secret',status='accepted')
        await db.create_workspace_memory('w',content='Suggested secret',status='suggested')
        await db.create_conversation('c',workspace_context={'mode':'selected','workspace_id':'w'})
        info=await db.build_workspace_memory_context(conversation_id='c')
        assert 'Use metric units.' in info['context'] and 'secret' not in info['context']
        conn=await db.get_db()
        try:
            await conn.execute('ALTER TABLE workspaces DROP COLUMN chat_defaults')
            await conn.execute('ALTER TABLE workspaces DROP COLUMN source_context_enabled')
            await conn.commit()
        finally:await conn.close()
        await db.init_db();await db.init_db()
        saved=await db.get_workspace('w')
        assert saved['source_context_enabled'] is True and saved['memory_enabled']==0 and saved['chat_defaults']=={}
        await db.update_workspace('w',memory_enabled=1)
        info=await db.build_workspace_memory_context(conversation_id='c')
        assert 'Accepted secret' in info['context'] and 'Suggested secret' not in info['context']
    asyncio.run(run())


async def seed_sources():
    await db.create_workspace('w','Project')
    await db.create_workspace('other','Other')
    await db.create_conversation('past','Earlier decisions',workspace_context={'mode':'selected','workspace_id':'w'})
    message=await db.add_message('past','assistant','cobaltnebula was assigned port 4872')
    await db.create_conversation('current',workspace_context={'mode':'selected','workspace_id':'w'})
    await db.add_message('current','assistant','cobaltnebula current duplicate')
    await db.create_conversation('unlinked')
    await db.add_message('unlinked','assistant','cobaltnebula isolated secret')
    await db.create_note('note',title='Port decision',content='cobaltnebula confirmed port 4872')
    await ws.link_source('w','note','note')
    await db.create_research_report('report',query='cobaltnebula',title='Cobalt report',status='completed')
    await db.update_research_report('report',report_markdown='cobaltnebula report uses port 4872')
    await ws.link_source('w','research','report')
    await db.add_artifact(artifact_id='artifact',conversation_id='past',filename='design.txt',url='/design.txt',content_text='cobaltnebula artifact uses port 4872')
    return message


def test_live_sources_exclusions_edits_unlinks_ownership_and_citations(isolated,monkeypatch):
    import rag
    monkeypatch.setattr(rag,'hybrid_query',AsyncMock(return_value=[]))
    monkeypatch.setattr(ws,'prepare_workspace_artifacts',AsyncMock())
    async def run():
        message=await seed_sources()
        async def retrieve():return await ws.retrieve('w','current','What port does cobaltnebula use?',[],10,6000)
        context,sources,partial=await retrieve()
        assert {s['kind'] for s in sources}=={'chat','note','research','artifact'}
        assert 'isolated secret' not in context and 'current duplicate' not in context
        chat=next(s for s in sources if s['kind']=='chat')
        assert chat['source_id']==str(message) and chat['parent_id']=='past' and 'assistant excerpt' in chat['location']
        assert 'parent=past' in chat['url'] and len(chat['revision'])==64
        await ws.set_source_preferences('w','note','note',{'enabled':False,'pinned':True})
        assert 'note' not in {s['kind'] for s in (await retrieve())[1]}
        assert (await ws.list_sources('w',pinned=True))['items'][0]['source_id']=='note'
        await db.update_message(message,content='cobaltnebula now uses port 5821')
        new=(await retrieve())[1]
        assert next(s for s in new if s['kind']=='chat')['revision']!=chat['revision']
        await ws.link_source('w','chat','past',remove=True)
        assert 'chat' not in {s['kind'] for s in (await retrieve())[1]}
        # Artifacts also have a direct workspace link assigned by delivery.
        await ws.link_source('w','artifact','artifact',remove=True)
        await ws.link_source('w','research','report',remove=True)
        assert not (await retrieve())[0]
        assert await db.get_research_report('report')
        user=await db.create_user('Other owner');token=db.set_current_user_id(user['id'])
        try:
            with pytest.raises(LookupError):await ws.retrieve('w','current','cobaltnebula',[],6,6000)
            with pytest.raises(LookupError):await ws.set_source_preferences('w','note','note',{'enabled':True})
        finally:db.reset_current_user_id(token)
    asyncio.run(run())


def test_picker_pages_all_types_and_does_not_duplicate_content(isolated):
    async def run():
        await seed_sources()
        for i in range(32):await db.create_note(f'n{i}',title=f'Archived decision {i}',content='kept')
        first=await ws.picker('w','note','Archived',limit=30)
        second=await ws.picker('w','note','Archived',limit=30,offset=30)
        assert len(first['items'])==30 and first['has_more'] and len(second['items'])==2
        for kind,identity in [('chat','unlinked'),('note','n31'),('research','report'),('artifact','artifact')]:
            await ws.link_source('other',kind,identity)
            assert await ws.source('other',kind,identity)
            await ws.link_source('other',kind,identity,remove=True)
            assert not await ws.source('other',kind,identity)
        assert (await db.get_note('n31'))['content']=='kept'
    asyncio.run(run())


def test_budgets_partial_failure_and_change_during_retrieval(isolated,monkeypatch):
    import rag,config
    monkeypatch.setattr(config,'RAG_RETRIEVAL_DEADLINE_SECONDS',.15)
    monkeypatch.setattr(ws,'prepare_workspace_artifacts',AsyncMock())
    async def run():
        await seed_sources();await db.create_kb('kb','Slow KB');await ws.link('w','knowledge','kb')
        cancelled=asyncio.Event()
        async def slow(*args,**kwargs):
            try:await asyncio.Event().wait()
            finally:cancelled.set()
        monkeypatch.setattr(rag,'hybrid_query',slow)
        start=time.monotonic()
        context,sources,partial=await ws.retrieve('w','current','cobaltnebula',[],3,120)
        assert time.monotonic()-start<.5 and partial and cancelled.is_set()
        assert context and len(context)<=120 and len(sources)<=3
        assert [s['n'] for s in sources]==list(range(1,len(sources)+1))
        assert all(f"[{s['n']}]" in context for s in sources)
        original=ws.library_passages
        async def changed(*args):
            result=await original(*args)
            await ws.link_source('w','note','note',remove=True)
            await db.delete_artifact('artifact')
            return result
        monkeypatch.setattr(ws,'library_passages',changed)
        assert not {'note','artifact'} & {s['kind'] for s in (await ws.retrieve('w','current','cobaltnebula',[],8,6000))[1]}
        await db.update_workspace('w',source_context_enabled=False)
        assert not (await ws.retrieve('w','current','cobaltnebula',[],8,6000))[0]
    asyncio.run(run())


def test_upload_originals_duplicate_names_unsupported_and_restart(isolated,monkeypatch):
    import indexing_jobs,rag
    from routes.context import route_context
    from routes.workspaces import upload
    from starlette.datastructures import UploadFile
    monkeypatch.setattr(route_context(),'track_bg',lambda coro:coro.close())
    async def run():
        await db.create_workspace('w','Project')
        first=await upload('w',UploadFile(io.BytesIO(b'first cobaltnebula'),filename='plan.txt'))
        second=await upload('w',UploadFile(io.BytesIO(b'second cobaltnebula'),filename='plan.txt'))
        unsupported=await upload('w',UploadFile(io.BytesIO(b'\x00\x01'),filename='binary.zip'))
        assert first['file_id']!=second['file_id'] and first['kb_id']==second['kb_id']
        files=await ws.source_files('w')
        assert len({f['filename'] for f in files})==3
        assert unsupported['status']=='unsupported'
        assert len((await ws.list_sources('w'))['items'])==3
        await indexing_jobs.reconcile()
        assert (await indexing_jobs.get(first['job_id']))['status']=='interrupted'
        monkeypatch.setattr(rag,'index_file',AsyncMock(return_value={'chunks':2,'warnings':['truncated']}))
        await indexing_jobs.retry(first['job_id'],lambda c:c.close())
        await indexing_jobs.execute(first['job_id'])
        assert (await indexing_jobs.get(first['job_id']))['status']=='partial'
        await db.delete_workspace('w')
        assert await db.get_kb(first['kb_id'])
        assert (await db.get_kb(first['kb_id']))['files']
    asyncio.run(run())


def test_artifact_file_revision_fencing_and_pdf_office_parsing(isolated,monkeypatch):
    import config,rag,indexing_jobs
    from routes.context import route_context
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfWriter
    monkeypatch.setattr(rag.ocr,'should_ocr',lambda *a:False)
    monkeypatch.setattr(route_context(),'track_bg',lambda c:c.close())
    async def run():
        await db.create_workspace('w','Project');await db.create_conversation('c',workspace_context={'mode':'selected','workspace_id':'w'})
        folder=Path(config.ARTIFACT_DIR);folder.mkdir()
        doc=Document();doc.add_paragraph('cobaltnebula office paragraph');doc.save(folder/'plan.docx')
        book=Workbook();book.active['A1']='cobaltnebula spreadsheet';book.save(folder/'plan.xlsx')
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='cobaltnebula slide';deck.save(folder/'plan.pptx')
        writer=PdfWriter();writer.add_blank_page(200,200);writer.write(folder/'empty.pdf')
        for ext in ('docx','xlsx','pptx'):
            assert 'cobaltnebula' in rag.parse_file(str(folder/f'plan.{ext}'),f'plan.{ext}')
        assert rag.parse_file(str(folder/'empty.pdf'),'empty.pdf')==''
        from pypdf.generic import NameObject,DictionaryObject,DecodedStreamObject
        writer=PdfWriter();page=writer.add_blank_page(200,200)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 10 150 Td (cobaltnebula PDF passage) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream);writer.write(folder/'text.pdf')
        assert 'cobaltnebula PDF passage' in rag.parse_file(str(folder/'text.pdf'),'text.pdf')
        await db.add_artifact(artifact_id='office',conversation_id='c',filename='plan.docx',url='/plan.docx',storage_path=str(folder/'plan.docx'))
        await ws.index_artifact('w','office',lambda c:c.close())
        mapped=(await ws.source_files('w'))[0]
        job=(await indexing_jobs.latest(mapped['kb_id'],mapped['filename']))['job_id']
        monkeypatch.setattr(rag,'index_file',AsyncMock(return_value={'chunks':1}))
        await indexing_jobs.execute(job)
        monkeypatch.setattr(rag,'hybrid_query',AsyncMock(return_value=[{'kb_id':mapped['kb_id'],'filename':mapped['filename'],'text':'cobaltnebula office paragraph','score':1}]))
        context,sources,_=await ws.retrieve('w','c','cobaltnebula',[],6,6000)
        assert context and sources[0]['kind']=='artifact' and sources[0]['source_id']=='office'
        await db.update_artifact_file_metadata('office',content_text='New revision')
        context,sources,_=await ws.retrieve('w','c','cobaltnebula',[],6,6000)
        assert not any(s.get('kb_id') for s in sources)
    asyncio.run(run())


@pytest.mark.parametrize('mode,ghost,expected',[('selected',False,True),('off',False,False),('auto',False,False),('selected',True,False)])
def test_workspace_to_answer_captured_context(isolated,monkeypatch,mode,ghost,expected):
    import rag,reasoning,context_inspector
    from agents import chat
    from main import ChatRequest
    monkeypatch.setattr(rag,'hybrid_query',AsyncMock(return_value=[]))
    monkeypatch.setattr(ws,'prepare_workspace_artifacts',AsyncMock())
    monkeypatch.setattr(chat,'_suggest_workspace_memories',AsyncMock())
    reasoning._cache.clear();payloads=[];emitted=[]
    def provider(request):
        if request.url.path=='/api/tags':return httpx.Response(200,json={'models':[{'name':'fixture','digest':'hub'}]})
        if request.url.path=='/api/show':return httpx.Response(200,json={'capabilities':['completion']})
        assert request.url.path=='/api/chat'
        payloads.append(json.loads(request.content))
        return httpx.Response(200,text=json.dumps({'message':{'role':'assistant','content':'Port 4872 [1].'},'done':True,'eval_count':5,'prompt_eval_count':30})+'\n')
    async def run():
        await seed_sources();await db.update_workspace('w',instructions='Always use metric units.',memory_enabled=0)
        await db.create_workspace_memory('w',content='Unaccepted dangerous memory',status='suggested')
        await db.add_conv_to_workspace('other','current')
        await db.update_conversation('current',workspace_context={'mode':mode,**({'workspace_id':'w'} if mode=='selected' else {})})
        async def emit(*args):emitted.append(args)
        req=ChatRequest(conversation_id='current',model='fixture',messages=[{'role':'user','content':'What port does cobaltnebula use?'}],use_memories=False,ephemeral=ghost,effort_mode='auto',reasoning_model='fixture')
        async with httpx.AsyncClient(transport=httpx.MockTransport(provider)) as http:
            events=[json.loads(e[5:]) async for e in chat.chat_stream_generate(req,http,SimpleNamespace(emit=emit),{}, {})]
        assert not [e for e in events if e.get('type')=='error']
        prompt=json.dumps(payloads)
        assert ('port 4872' in prompt)==expected
        assert ('Always use metric units.' in prompt)==expected
        assert 'Unaccepted dangerous memory' not in prompt and 'isolated secret' not in prompt
        if expected:
            assert 'not accepted memories' in prompt
            conv=await db.get_conversation('current');record=(await context_inspector.inspect(conv['messages'][-1]['id']))['rounds'][0]
            assert {s['kind'] for s in record['retrieved_sources']}=={'chat','note','research','artifact'}
            assert record['retrieved_context'] in payloads[0]['messages'][0]['content']
            assert not record['memories']['workspace']['ids']
    asyncio.run(run())


def test_hub_routes_policy_and_owner_checks(isolated,monkeypatch):
    import main
    from routes.context import route_context
    monkeypatch.setattr(route_context(),'track_bg',lambda c:c.close())
    async def run():
        await seed_sources()
        await db.create_workspace_memory('w',content='Pending decision')
        other=await db.create_user('Other owner')
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='http://fixture') as http:
            for path in ['defaults','sources','overview','picker/note']:
                response=await http.get('/api/workspaces/w/'+path)
                assert response.status_code==200,response.text
                foreign=await http.get('/api/workspaces/w/'+path,headers={'X-HyprChat-User':other['id']})
                assert foreign.status_code==404,foreign.text
            overview=(await http.get('/api/workspaces/w/overview')).json()
            assert overview['pending_memories']==1
            assert (await http.patch('/api/workspaces/w/sources/note/note',json={'enabled':False})).status_code==200
            assert (await http.patch('/api/workspaces/w/sources/note/note',json={'enabled':'false'})).status_code==400
            uploaded=await http.post('/api/workspaces/w/upload',files={'file':('same.txt',b'cobaltnebula','text/plain')})
            assert uploaded.status_code==200,uploaded.text
            assert uploaded.json()['job_id']
            file=uploaded.json()
            raw=await http.get(f"/api/knowledge-bases/{file['kb_id']}/files/{file['file_id']}/raw")
            assert raw.status_code==200 and raw.content==b'cobaltnebula'
            assert 'same.txt' in raw.headers['content-disposition']
            preview=await http.get(f"/api/knowledge-bases/{file['kb_id']}/files/{file['file_id']}/preview")
            assert preview.status_code==200 and preview.json()['filename']=='same.txt'
            assert (await http.post('/api/workspaces/other/sources/note/note/link')).status_code==200
            assert (await http.delete('/api/workspaces/other/sources/note/note/link')).status_code==200
    asyncio.run(run())


def test_embedding_failure_is_durable_and_progress_is_real(isolated,monkeypatch):
    import rag,indexing_jobs,job_registry
    monkeypatch.setattr(rag,'embed_texts',AsyncMock(return_value=[None]))
    async def run():
        await db.create_kb('kb','Documents')
        original=isolated/'original.txt';original.write_text('cobaltnebula source passage')
        await db.add_kb_file('kb','original.txt',str(original),original.stat().st_size,'text/plain')
        identity=await indexing_jobs.submit('kb','original.txt',str(original))
        await indexing_jobs.execute(identity)
        job=await indexing_jobs.get(identity)
        assert job['status']=='failed' and 'embeddings' in job['error']
        assert original.read_text()=='cobaltnebula source passage'
        assert not (await ws.knowledge_status(await db.get_kb('kb')))['ready']
        token=indexing_jobs.current_job.set(identity)
        try:
            await indexing_jobs.retry(identity,lambda c:c.close())
            conn=await db.get_db()
            await conn.execute("UPDATE indexing_jobs SET status='running' WHERE id=?",(identity,));await conn.commit();await conn.close()
            await indexing_jobs.report_progress('embedding',64,128)
            progress=await job_registry.get('indexing:'+identity)
            assert progress['phase']=='embedding' and progress['details']['completed']==64
        finally:indexing_jobs.current_job.reset(token)
    asyncio.run(run())
