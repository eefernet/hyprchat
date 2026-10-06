"""Real SQLite behavior with a deterministic document worker transport."""
import asyncio
import base64
import io
import json
import sys
import types
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
pytest.importorskip("fastapi")
pytest.importorskip("docx")
pytest.importorskip("lxml")
from fastapi import HTTPException, UploadFile
from .optional_deps import load_route_module
import database as db
import documents
from document_runtime import execute, create_document


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DATABASE_PATH", str(tmp_path / "db.sqlite"))
    monkeypatch.setattr(documents.config, "SANDBOX_OUTPUTS_DIR", str(tmp_path / "outputs"))
    settings = tmp_path / "settings.json"
    settings.write_text('{"document_tools_enabled":true}')
    monkeypatch.setattr(documents.config, "SETTINGS_PATH", str(settings))
    token = db.set_current_user_id(db.DEFAULT_USER_ID)
    asyncio.run(db.init_db())
    asyncio.run(db.create_conversation("conv-docs", "Document test"))
    route = load_route_module(monkeypatch, "chat_files")
    monkeypatch.setitem(sys.modules, "routes.chat_files", route)
    parent = types.ModuleType("routes")
    parent.__path__ = []
    monkeypatch.setitem(sys.modules, "routes", parent)
    states = {}
    async def upload(http, path, data):
        relative = path.removeprefix(documents.JOB_ROOT + "/")
        local = tmp_path / "worker" / relative
        local.parent.mkdir(parents=True, exist_ok=True)
        local.write_bytes(data)
    monkeypatch.setattr(route, "upload_bytes_to_codebox", upload)
    async def worker(http, command, *args):
        rid = args[0] if args else ""
        directory = tmp_path / "worker" / rid
        if command == "start":
            request = json.loads((directory / "request.json").read_text())
            try:
                result = execute(request, directory)
                for output in result["outputs"]:
                    output["size_bytes"] = (directory / output["path"]).stat().st_size
                states[rid] = {"status": "succeeded", **result}
            except Exception as e:
                states[rid] = {"status": "failed", "error": str(e)}
            return {"status": "queued"}
        if command == "status":
            return states[rid]
        if command == "chunk":
            data = (directory / args[1]).read_bytes()[int(args[2]):int(args[2]) + 192 * 1024]
            return {"data": base64.b64encode(data).decode()}
        return {"status": "cancelled"}
    monkeypatch.setattr(documents, "worker", worker)
    async def emit(*args):
        pass
    monkeypatch.setattr(documents, "_events", types.SimpleNamespace(emit=emit))
    yield tmp_path, settings, worker
    db.reset_current_user_id(token)


def upload(path):
    return UploadFile(filename=path.name, file=io.BytesIO(path.read_bytes()))


def test_upload_read_edit_revision_and_duplicate_names(env):
    tmp, _, _ = env
    path = tmp / "report.docx"
    create_document(path, {"blocks": [{"text": "Original"}]})
    async def scenario():
        first = await documents.upload_document(None, upload(path), "conv-docs")
        second = await documents.upload_document(None, upload(path), "conv-docs")
        assert first["artifact_id"] != second["artifact_id"]
        read = await documents.submit(None, {"action": "read", "artifact_id": first["artifact_id"], "conversation_id": "conv-docs"}, wait=True)
        assert read["status"] == "succeeded", read
        target = read["report"]["items"][0]["target"]
        edited = await documents.submit(None, {"action": "edit", "artifact_id": first["artifact_id"], "expected_sha256": read["source_sha256"],
                                               "operations": [{"op": "set_text", "target": target, "text": "Revised"}], "conversation_id": "conv-docs"}, wait=True)
        assert edited["status"] == "succeeded", edited
        revision = await db.get_artifact(edited["artifacts"][0]["artifact_id"])
        assert revision["supersedes_artifact_id"] == first["artifact_id"]
        original, original_path = await documents.owned_source(first["artifact_id"])
        assert original_path.read_bytes() == path.read_bytes()
        assert revision["storage_path"] != original["storage_path"]
        assert (await db.get_run(edited["run_id"]))["status"] == "succeeded"
    asyncio.run(scenario())


def test_owned_artifacts_and_stale_hash(env):
    tmp, _, _ = env
    path = tmp / "report.docx"
    create_document(path, {"blocks": [{"text": "Original"}]})
    async def scenario():
        first = await documents.upload_document(None, upload(path), "conv-docs")
        with pytest.raises(HTTPException) as stale:
            await documents.submit(None, {"action": "edit", "artifact_id": first["artifact_id"], "expected_sha256": "bad", "conversation_id": "conv-docs"}, wait=True)
        assert stale.value.status_code == 409
        token = db.set_current_user_id("another-user")
        try:
            with pytest.raises(HTTPException) as missing:
                await documents.owned_source(first["artifact_id"])
            assert missing.value.status_code == 404
        finally:
            db.reset_current_user_id(token)
    asyncio.run(scenario())


def test_disabled_and_invalid_binary_upload(env):
    _, settings, _ = env
    async def scenario():
        with pytest.raises(HTTPException) as malformed:
            await documents.upload_document(None, UploadFile(filename="bad.docx", file=io.BytesIO(b"plain text")), "conv-docs")
        assert malformed.value.status_code == 422
        settings.write_text('{"document_tools_enabled":false}')
        with pytest.raises(HTTPException) as disabled:
            await documents.submit(None, {"action": "create", "conversation_id": "conv-docs", "format": "docx"}, wait=True)
        assert disabled.value.status_code == 503
        assert await db.list_artifacts() == []
    asyncio.run(scenario())


def test_cancel_during_transfer_publishes_nothing(env, monkeypatch):
    _, _, actual = env
    async def worker(http, command, *args):
        if command == "chunk":
            documents.cancel_registry.signal(args[0])
        return await actual(http, command, *args)
    monkeypatch.setattr(documents, "worker", worker)
    async def scenario():
        result = await documents.submit(None, {"action": "create", "format": "docx", "content": {"blocks": [{"text": "Never publish"}]}, "conversation_id": "conv-docs"}, wait=True)
        assert result["status"] == "cancelled", result
        assert await db.list_artifacts() == []
    asyncio.run(scenario())


def test_database_failure_rolls_back_artifacts_and_files(env, monkeypatch):
    tmp, _, _ = env
    original = db.add_artifact
    async def broken(**kwargs):
        await original(**kwargs)
        raise RuntimeError("Injected publication failure")
    monkeypatch.setattr(db, "add_artifact", broken)
    async def scenario():
        result = await documents.submit(None, {"action": "create", "format": "docx", "content": {"blocks": [{"text": "Never publish"}]}, "conversation_id": "conv-docs"}, wait=True)
        assert result["status"] == "failed", result
        assert await db.list_artifacts() == []
        assert list((tmp / "outputs").glob("hc-document-*")) == []
    asyncio.run(scenario())


def test_worker_failure_retains_no_success_artifact(env, monkeypatch):
    async def unavailable(*args):
        raise RuntimeError("Codebox unavailable")
    monkeypatch.setattr(documents, "worker", unavailable)
    result = asyncio.run(documents.submit(None, {"action": "create", "format": "docx", "content": {"blocks": [{"text": "Test"}]}, "conversation_id": "conv-docs"}, wait=True))
    assert result["status"] == "failed"
    assert "unavailable" in result["error"]
    assert asyncio.run(db.list_artifacts()) == []


def test_shutdown_cancels_worker_and_publishes_nothing(env, monkeypatch):
    _, _, actual = env
    async def scenario():
        started = asyncio.Event()
        cancelled = []
        async def worker(http, command, *args):
            if command == 'start':
                started.set()
                await asyncio.Event().wait()
            if command == 'cancel':
                cancelled.append(args[0])
            return await actual(http, command, *args)
        monkeypatch.setattr(documents, 'worker', worker)
        job = await documents.submit(None, {'action': 'create', 'format': 'docx', 'content': {'blocks': [{'text': 'Test'}]}, 'conversation_id': 'conv-docs'})
        await asyncio.wait_for(started.wait(), 2)
        await documents.shutdown()
        assert cancelled == [job['run_id']]
        assert (await db.get_run(job['run_id']))['status'] == 'cancelled'
        assert await db.list_artifacts() == []
    asyncio.run(scenario())


def test_tool_bounds_one_large_paragraph_without_invalid_json(env):
    tmp, _, _ = env
    source = tmp / 'long.docx'
    create_document(source, {'blocks': [{'text': 'x' * 25000}]})
    async def scenario():
        artifact = await documents.upload_document(None, upload(source), 'conv-docs')
        response = json.loads(await documents.tool('document_read', {'artifact_id': artifact['artifact_id']}, http=None, conv_id='conv-docs'))
        paragraph = response['report']['items'][0]
        assert len(paragraph['text']) == 12000
        assert paragraph['text_truncated'] is True
        assert response['status'] == 'succeeded'
    asyncio.run(scenario())


@pytest.mark.parametrize('content', ['{"blocks": [[broken]}', {'blocks': [{'type': 'table', 'rows': [['A', 'B'], ['short']]}]}, {}, {'blocks': 'text'}])
def test_invalid_content_rejected_before_worker_or_run(env, monkeypatch, content):
    async def forbidden(*args):
        raise AssertionError('Invalid content must not reach worker')
    monkeypatch.setattr(documents, 'worker', forbidden)
    async def scenario():
        result = json.loads(await documents.tool('document_create', {'format': 'docx', 'content': content}, http=None, conv_id='conv-docs'))
        assert result['status'] == 'failed'
        assert result['error_code'] == 'invalid_arguments'
        assert 'content' in result['error'] and 'Example content' in result['error']
        assert await db.get_runs_by_conversation('conv-docs') == []
        assert await db.list_artifacts() == []
    asyncio.run(scenario())


def test_encoded_json_content_creates_real_word_with_table(env):
    from docx import Document
    async def scenario():
        result = json.loads(await documents.tool('document_create', {'format': 'docx', 'content': json.dumps({'blocks': [{'type': 'table', 'rows': [['Name', 'Value'], ['A', '10']]}]})}, http=None, conv_id='conv-docs'))
        assert result['status'] == 'succeeded', result
        _, path = await documents.owned_source(result['artifacts'][0]['artifact_id'])
        assert Document(path).tables[0].cell(1, 1).text == '10'
    asyncio.run(scenario())


def test_markdown_export_preserves_sections_tables_and_code(env):
    from docx import Document
    markdown = '# Full report\nIntro with **bold** text.\n\n| Name | Value |\n| --- | --- |\n| A \\| B | 10 |\n\n## Later section\nDo not omit this.\n```python\n# literal code\nx = 1 | 2\n```\nFinal sentence.'
    async def scenario():
        result = json.loads(await documents.tool('document_create', {'format': 'docx', 'markdown': markdown}, http=None, conv_id='conv-docs'))
        assert result['status'] == 'succeeded', result
        _, path = await documents.owned_source(result['artifacts'][0]['artifact_id'])
        doc = Document(path)
        assert doc.paragraphs[0].text == 'Full report'
        assert doc.paragraphs[0].style.name == 'Heading 1'
        assert any(r.text == 'bold' and r.bold for p in doc.paragraphs for r in p.runs)
        assert doc.tables[0].cell(1, 0).text == 'A | B'
        text = '\n'.join(p.text for p in doc.paragraphs)
        assert 'Later section' in text and '# literal code\nx = 1 | 2' in text
        assert text.endswith('Final sentence.')
    asyncio.run(scenario())


@pytest.mark.parametrize('args', [
    {'format': 'pptx', 'markdown': '# Report'},
    {'format': 'docx', 'markdown': '   '},
    {'format': 'docx', 'markdown': '# Report', 'content': {'blocks': [{'text': 'Conflicting'}]}},
])
def test_invalid_markdown_request_has_no_worker_run(env, args):
    async def scenario():
        result = json.loads(await documents.tool('document_create', args, http=None, conv_id='conv-docs'))
        assert result['error_code'] == 'invalid_arguments'
        assert await db.get_runs_by_conversation('conv-docs') == []
    asyncio.run(scenario())


def test_export_content_text_uses_only_exported_sheet(env):
    from openpyxl import Workbook
    tmp, _, _ = env
    path = tmp / 'two-sheets.xlsx'
    book = Workbook()
    book.active.title = 'Selected'
    book.active.append(['CSV_ONLY_VALUE'])
    book.create_sheet('Private').append(['OTHER_SHEET_SENTINEL'])
    book.save(path)
    async def scenario():
        source = await documents.upload_document(None, upload(path), 'conv-docs')
        result = await documents.submit(None, {'action': 'convert', 'artifact_id': source['artifact_id'],
            'conversation_id': 'conv-docs', 'format': 'csv', 'sheet': 'Selected'}, wait=True)
        assert result['status'] == 'succeeded', result
        exported = await db.get_artifact(result['artifacts'][0]['artifact_id'])
        assert 'CSV_ONLY_VALUE' in exported['content_text']
        assert 'OTHER_SHEET_SENTINEL' not in exported['content_text']
    asyncio.run(scenario())


def test_missing_cached_pdf_has_no_render_url(env):
    async def scenario():
        preview = await documents.store_file(b'%PDF-stub', 'preview.pdf', conversation_id='conv-docs', role='preview')
        source = {'url': '/original', 'metadata': {'document_preview_id': preview['id']}}
        assert (await documents.preview_payload(source))['render_url'] == preview['url']
        Path(preview['storage_path']).unlink()
        assert (await documents.preview_payload(source))['render_url'] is None
    asyncio.run(scenario())


def test_cancel_failure_is_retryable_and_late_stop_keeps_success(env, monkeypatch):
    _, _, actual = env
    async def unavailable(*args):
        raise ConnectionError('offline')
    async def scenario():
        await db.create_run('doc-stop', 'conv-docs', 'documents', status='running')
        monkeypatch.setattr(documents, 'worker', unavailable)
        response = await documents.cancel_run(None, 'doc-stop')
        assert response['status'] == 'cancelling'
        assert 'unconfirmed' in response['error']
        assert (await db.get_run('doc-stop'))['ended_at'] is None
        monkeypatch.setattr(documents, 'worker', actual)
        assert (await documents.cancel_run(None, 'doc-stop'))['status'] == 'cancelled'
        await db.update_run('doc-stop', status='succeeded', ended=True)
        monkeypatch.setattr(documents, 'worker', unavailable)
        assert (await documents.cancel_run(None, 'doc-stop'))['status'] == 'succeeded'
    asyncio.run(scenario())


def test_stop_waits_for_worker_terminal_state(env, monkeypatch):
    async def scenario():
        await db.create_run('doc-stop', 'conv-docs', 'documents', status='running')
        polled = asyncio.Event()
        release = asyncio.Event()
        async def worker(http, command, *args):
            if command == 'cancel':
                return {'status': 'running'}
            polled.set()
            await release.wait()
            return {'status': 'cancelled'}
        monkeypatch.setattr(documents, 'worker', worker)
        task = asyncio.create_task(documents.cancel_run(None, 'doc-stop'))
        await asyncio.wait_for(polled.wait(), 2)
        assert not task.done()
        assert (await db.get_run('doc-stop'))['status'] == 'cancelling'
        release.set()
        assert (await task)['status'] == 'cancelled'
    asyncio.run(scenario())


def test_chat_wait_jobs_are_tracked_and_settled_on_shutdown(env, monkeypatch):
    _, _, actual = env
    async def scenario():
        started = asyncio.Event()
        async def worker(http, command, *args):
            if command == 'start':
                started.set()
                await asyncio.Event().wait()
            return await actual(http, command, *args)
        monkeypatch.setattr(documents, 'worker', worker)
        caller = asyncio.create_task(documents.submit(None, {'action': 'create', 'format': 'docx',
            'content': {'blocks': [{'text': 'Shutdown test'}]}, 'conversation_id': 'conv-docs'}, wait=True))
        await asyncio.wait_for(started.wait(), 2)
        assert len(documents._tasks) == 1
        rid = next(iter(documents._tasks))
        await documents.shutdown()
        await asyncio.gather(caller, return_exceptions=True)
        assert not documents._tasks
        assert (await db.get_run(rid))['status'] == 'cancelled'
        assert await db.list_artifacts() == []
    asyncio.run(scenario())


def test_startup_reaps_unconfirmed_stops_without_claiming_stopped(env):
    async def scenario():
        await db.create_run('doc-stop', 'conv-docs', 'documents', status='cancelling')
        await db.reap_stale_runs()
        row = await db.get_run('doc-stop')
        assert row['status'] == 'failed'
        assert 'restart' in row['result_envelope']['summary']
    asyncio.run(scenario())
